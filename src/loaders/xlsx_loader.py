"""XLSX loader — per-sheet, 50-row blocks with repeated header."""
import logging
from pathlib import Path

from langchain_core.documents import Document

log = logging.getLogger("rag")

_BLOCK_SIZE = 50


def _format_row(row: tuple) -> str:
    return "\t".join("" if c is None else str(c) for c in row)


def load(path: Path) -> list[Document]:
    try:
        import openpyxl
    except ImportError as e:
        raise ImportError("openpyxl is required: pip install openpyxl") from e

    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as e:
        raise ValueError(f"Invalid XLSX: {path.name}") from e

    docs: list[Document] = []

    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header, data_rows = rows[:1], rows[1:]
        if not header or not data_rows:
            continue

        total_rows = len(data_rows)
        if total_rows > 10000:
            log.warning(
                "[xlsx_loader] %s#%s has %d rows, expect slow ingest",
                path.name, sheet.title, total_rows,
            )

        filename = f"{path.name}#{sheet.title}"
        header_line = _format_row(header[0])

        for block_start in range(0, total_rows, _BLOCK_SIZE):
            block = data_rows[block_start: block_start + _BLOCK_SIZE]
            body_lines = "\n".join(_format_row(row) for row in block)

            docs.append(Document(
                page_content=header_line + "\n" + body_lines,
                metadata={
                    "filename": filename,
                    "page": 0,
                    "sheet": sheet.title,
                    "row_range": f"{block_start + 1}-{block_start + len(block)}",
                },
            ))

    log.info("[xlsx_loader] %s → %d document(s)", path.name, len(docs))
    return docs
