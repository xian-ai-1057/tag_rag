"""Unit tests for src/loaders/xlsx_loader.py."""
import importlib.util
from pathlib import Path

import pytest
from langchain_core.documents import Document

from src.loaders import load_file

_REPO_ROOT = Path(__file__).resolve().parents[2]
_SCHEMA_PATH = _REPO_ROOT / "specs" / "102-multi-format-loaders" / "contracts" / "loaded_doc_metadata_schema.py"


def _load_metadata_schema():
    spec = importlib.util.spec_from_file_location("loaded_doc_metadata_schema", _SCHEMA_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.LoadedDocMetadata


def _build_xlsx_2sheets(tmp_path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["Name", "Value"])
    ws1.append(["Alpha", 100])
    ws1.append(["Beta", 200])

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Item", "Count"])
    ws2.append(["X", 10])
    ws2.append(["Y", 20])

    p = tmp_path / "tiny.xlsx"
    wb.save(str(p))
    return p


def _build_xlsx_120rows(tmp_path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["ID", "Score"])
    for i in range(1, 121):
        ws.append([i, i * 10])

    p = tmp_path / "large.xlsx"
    wb.save(str(p))
    return p


@pytest.fixture
def xlsx_path(tmp_path):
    return _build_xlsx_2sheets(tmp_path)


def test_xlsx_returns_documents(xlsx_path):
    docs = load_file(xlsx_path)
    assert len(docs) >= 2
    assert all(isinstance(d, Document) for d in docs)


def test_xlsx_multisheet_filenames_differ(xlsx_path):
    docs = load_file(xlsx_path)
    filenames = {d.metadata["filename"] for d in docs}
    assert len(filenames) == 2, f"Expected 2 distinct filenames, got: {filenames}"
    for fn in filenames:
        assert "#" in fn, f"filename missing '#': {fn}"
    assert "tiny.xlsx#Sheet1" in filenames
    assert "tiny.xlsx#Sheet2" in filenames


def test_xlsx_different_sheet_file_ids_differ(xlsx_path):
    from src.chunking import chunk_documents

    docs = load_file(xlsx_path)
    chunks = chunk_documents(docs)
    file_ids = {c.metadata["file_id"] for c in chunks}
    assert len(file_ids) == 2, f"Expected 2 distinct file_ids, got: {file_ids}"


def test_xlsx_metadata_schema(xlsx_path):
    LoadedDocMetadata = _load_metadata_schema()
    docs = load_file(xlsx_path)
    for d in docs:
        LoadedDocMetadata.model_validate(d.metadata)


def test_xlsx_50row_blocking(tmp_path):
    p = _build_xlsx_120rows(tmp_path)
    docs = load_file(p)
    assert len(docs) == 3, f"Expected 3 docs for 120 data rows, got {len(docs)}"
    row_ranges = [d.metadata["row_range"] for d in docs]
    assert "1-50" in row_ranges
    assert "51-100" in row_ranges
    assert "101-120" in row_ranges


def test_xlsx_header_repeated_in_all_blocks(tmp_path):
    p = _build_xlsx_120rows(tmp_path)
    docs = load_file(p)
    for d in docs:
        first_line = d.page_content.splitlines()[0]
        assert "ID" in first_line and "Score" in first_line, (
            f"Header not repeated in block with row_range={d.metadata['row_range']}"
        )


def test_xlsx_sheet_metadata(xlsx_path):
    docs = load_file(xlsx_path)
    for d in docs:
        assert d.metadata["sheet"] is not None
        assert d.metadata["page"] == 0
        assert d.metadata["row_range"] is not None
