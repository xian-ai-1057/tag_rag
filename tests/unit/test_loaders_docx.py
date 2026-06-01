"""Unit tests for src/loaders/docx_loader.py."""
import importlib.util
from pathlib import Path

import pytest
from langchain_core.documents import Document

from src.loaders import load_file

_REPO_ROOT = Path(__file__).resolve().parents[2]
_SCHEMA_PATH = _REPO_ROOT / "specs" / "102-multi-format-loaders" / "contracts" / "loaded_doc_metadata_schema.py"


def _load_metadata_schema():
    spec = importlib.util.spec_from_file_location("loaded_doc_metadata_schema", _SCHEMA_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.LoadedDocMetadata


def _build_docx(tmp_path: Path) -> Path:
    from docx import Document as DocxDoc

    doc = DocxDoc()
    doc.add_heading("信用卡業務", level=1)
    doc.add_paragraph("本段落說明信用卡業務第一季度表現。")
    doc.add_paragraph("第二段文字補充細節。")
    p = tmp_path / "tiny.docx"
    doc.save(str(p))
    return p


@pytest.fixture
def docx_path(tmp_path):
    return _build_docx(tmp_path)


def test_docx_returns_documents(docx_path):
    docs = load_file(docx_path)
    assert len(docs) >= 1
    assert all(isinstance(d, Document) for d in docs)


def test_docx_section_captured(docx_path):
    docs = load_file(docx_path)
    sections = [d.metadata.get("section") for d in docs]
    assert any(s is not None for s in sections), "Expected at least one doc with non-None section"
    assert any(s == "信用卡業務" for s in sections if s), "Expected section == '信用卡業務'"


def test_docx_metadata_schema(docx_path):
    LoadedDocMetadata = _load_metadata_schema()
    docs = load_file(docx_path)
    for d in docs:
        LoadedDocMetadata.model_validate(d.metadata)


def test_docx_filename_is_basename(docx_path):
    docs = load_file(docx_path)
    for d in docs:
        assert d.metadata["filename"] == docx_path.name


def test_docx_no_heading_section_none(tmp_path):
    from docx import Document as DocxDoc

    doc = DocxDoc()
    doc.add_paragraph("無標題段落一。")
    doc.add_paragraph("無標題段落二。")
    p = tmp_path / "no_heading.docx"
    doc.save(str(p))

    docs = load_file(p)
    assert len(docs) >= 1
    for d in docs:
        assert d.metadata["section"] is None


def test_docx_unsupported_ext_raises():
    with pytest.raises(ValueError, match=r"\.(docx|xlsx|html)"):
        load_file(Path("foo.pptx"))
